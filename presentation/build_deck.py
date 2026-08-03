#!/usr/bin/env python3
"""
Aegis AI — investor / final-year pitch deck generator.

Design principle (matches the platform's trust-integrity work): capabilities are
split honestly into BUILT TODAY vs VISION / ROADMAP. Nothing unbuilt is claimed
as deployed. Every slide carries speaker notes with an objective, a spoken
script, and — where useful — anticipated jury questions with answers.

Run:  .venv/bin/python build_deck.py   ->  AegisAI_Pitch.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY      = RGBColor(0x0A, 0x0F, 0x20)   # page background
NAVY2     = RGBColor(0x10, 0x18, 0x30)   # panels
NAVY3     = RGBColor(0x16, 0x20, 0x40)   # raised panels
ROYAL     = RGBColor(0x3B, 0x6E, 0xF5)   # brand
ROYAL_DK  = RGBColor(0x24, 0x4A, 0xC4)
CYAN      = RGBColor(0x2D, 0xD4, 0xE0)   # accent
WHITE     = RGBColor(0xF5, 0xF8, 0xFF)
MUTE      = RGBColor(0x9A, 0xA8, 0xC8)   # muted text
MUTE2     = RGBColor(0x6C, 0x7A, 0x9C)   # captions
GREEN     = RGBColor(0x35, 0xD0, 0x8A)   # "built today"
AMBER     = RGBColor(0xF4, 0xB4, 0x40)   # "vision"
LINE      = RGBColor(0x2A, 0x36, 0x58)

FONT = "Segoe UI"
FONT_L = "Segoe UI Light"
FONT_SB = "Segoe UI Semibold"

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

_num = [0]


# ── Low-level helpers ────────────────────────────────────────────────────────
def _noline(shape):
    shape.line.fill.background()


def solid(shape, rgb, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    if line is None:
        _noline(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def rect(slide, x, y, w, h, rgb, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    return solid(sp, rgb, line)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def run(p, text, size, color=WHITE, bold=False, font=FONT, italic=False, spacing=None):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return r


def para(tf, first=False, space_before=0, space_after=6, align=PP_ALIGN.LEFT, level=0, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.alignment = align
    p.level = level
    if line is not None:
        p.line_spacing = line
    return p


def bg(slide, rgb=NAVY):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(section_tag=""):
    _num[0] += 1
    s = prs.slides.add_slide(BLANK)
    bg(s)
    return s


def footer(slide, section=""):
    _, tf = textbox(slide, Inches(0.7), Inches(7.02), Inches(9), Inches(0.35))
    p = para(tf, first=True, space_after=0)
    run(p, "AEGIS AI", 9, CYAN, bold=True)
    run(p, "   ·   The AI Operating System for Insurance", 9, MUTE2)
    if section:
        run(p, f"    ·   {section}", 9, MUTE2)
    _, tf2 = textbox(slide, Inches(12.2), Inches(7.02), Inches(0.9), Inches(0.35))
    pr = para(tf2, first=True, space_after=0, align=PP_ALIGN.RIGHT)
    run(pr, f"{_num[0]:02d}", 9, MUTE2, bold=True)


def header(slide, title, subtitle=None, section="", kicker=None):
    """Standard slide header: accent bar, kicker, title, subtitle."""
    rect(slide, Inches(0.7), Inches(0.6), Inches(0.09), Inches(0.62), CYAN)
    x = Inches(0.95)
    _, tf = textbox(slide, x, Inches(0.5), Inches(11.6), Inches(1.15))
    if kicker:
        pk = para(tf, first=True, space_after=2)
        run(pk, kicker.upper(), 11, CYAN, bold=True)
        pt = para(tf, space_after=0)
    else:
        pt = para(tf, first=True, space_after=0)
    run(pt, title, 29, WHITE, bold=True, font=FONT_SB)
    if subtitle:
        ps = para(tf, space_before=2, space_after=0)
        run(ps, subtitle, 14, MUTE)
    footer(slide, section)


def caption(slide, text, y=Inches(6.55)):
    _, tf = textbox(slide, Inches(0.95), y, Inches(11.5), Inches(0.4))
    p = para(tf, first=True, space_after=0)
    run(p, "▣  Suggested visual:  ", 10.5, CYAN, bold=True)
    run(p, text, 10.5, MUTE)


def pill(slide, x, y, text, fg, bg_rgb, w=Inches(1.6)):
    sp = rect(slide, x, y, w, Inches(0.34), bg_rgb, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame
    tf.word_wrap = True
    p = para(tf, first=True, space_after=0, align=PP_ALIGN.CENTER)
    run(p, text, 10, fg, bold=True)
    return sp


# ── Content block renderer ───────────────────────────────────────────────────
def blocks(slide, items, x=Inches(0.95), y=Inches(1.9), w=Inches(11.6), h=Inches(4.5),
           size=14, gap=7):
    """items: list of tuples.
       ('H', text)        section heading (cyan)
       ('B', text)        bullet
       ('B2', text)       sub-bullet
       ('K', key, value)  key : value
       ('P', text)        plain paragraph
       ('S',)             small spacer
    """
    _, tf = textbox(slide, x, y, w, h)
    first = True
    for it in items:
        kind = it[0]
        if kind == 'H':
            p = para(tf, first=first, space_before=6, space_after=3)
            run(p, it[1].upper(), size - 1.5, CYAN, bold=True, font=FONT_SB)
        elif kind == 'B':
            p = para(tf, first=first, space_after=gap, line=1.05)
            run(p, "▸  ", size, ROYAL, bold=True)
            run(p, it[1], size, WHITE)
        elif kind == 'B2':
            p = para(tf, first=first, space_after=gap - 2, level=1)
            run(p, "–  ", size - 1, MUTE)
            run(p, it[1], size - 1, MUTE)
        elif kind == 'K':
            p = para(tf, first=first, space_after=gap, line=1.05)
            run(p, it[1] + "   ", size, CYAN, bold=True)
            run(p, it[2], size, WHITE)
        elif kind == 'P':
            p = para(tf, first=first, space_after=gap + 2, line=1.1)
            run(p, it[1], size, MUTE)
        elif kind == 'S':
            p = para(tf, first=first, space_after=4)
            run(p, "", size)
        first = False
    return tf


def two_col(slide, left, right, y=Inches(1.95), h=Inches(4.4)):
    blocks(slide, left, x=Inches(0.95), y=y, w=Inches(5.7), h=h)
    blocks(slide, right, x=Inches(7.0), y=y, w=Inches(5.5), h=h)


def panel(slide, x, y, w, h, rgb=NAVY2, line=LINE):
    return rect(slide, x, y, w, h, rgb, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def N(objective, script, qa=None):
    s = f"🎯 OBJECTIVE\n{objective}\n\n🎤 SCRIPT\n{script}"
    if qa:
        s += "\n\n❓ ANTICIPATED JURY Q&A"
        for q, a in qa:
            s += f"\nQ: {q}\nA: {a}\n"
    return s


# =============================================================================
#  SECTION A — COVER & FRAMING
# =============================================================================

# 1 — COVER
s = new_slide()
rect(s, 0, 0, SW, SH, NAVY)
# accent glow bars
rect(s, 0, Inches(6.9), SW, Inches(0.6), NAVY2)
rect(s, Inches(0.9), Inches(2.15), Inches(1.4), Inches(0.11), CYAN)
_, tf = textbox(s, Inches(0.9), Inches(2.35), Inches(11.5), Inches(2.6))
p = para(tf, first=True, space_after=0)
run(p, "AEGIS ", 66, WHITE, bold=True, font=FONT_SB)
run(p, "AI", 66, CYAN, bold=True, font=FONT_SB)
p2 = para(tf, space_before=6, space_after=0)
run(p2, "The Intelligent AI Infrastructure Platform for the Future of Insurance", 22, MUTE)
_, tf3 = textbox(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.0))
p3 = para(tf3, first=True, space_after=0)
run(p3, "Multi-Agent AI   ·   Explainable AI   ·   Decision Intelligence   ·   Enterprise API Platform", 14, CYAN, bold=True)
_, tf4 = textbox(s, Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.6))
p4 = para(tf4, first=True, space_after=0)
run(p4, "Final-Year Engineering Project  ×  AI Startup     |     Sri Nivaspanneer", 13, MUTE2)
notes(s, N(
    "Open with a bold, single identity: Aegis AI is infrastructure, not a chatbot.",
    "Good [morning/afternoon]. I'm Sri, and this is Aegis AI — an intelligent AI "
    "infrastructure platform for insurance. Not a chatbot, not a single model — an "
    "AI operating system that any insurer can plug into through APIs. Today I'll show "
    "you what we've built, the research behind it, and where this becomes a company.",
    [("Is this a product or a research project?",
      "Both — it is engineered as a working platform and grounded in current insurance-AI "
      "research. That dual nature is deliberate: a final-year project with startup ambition.")]))

# 2 — WHAT IS AEGIS (one-liner)
s = new_slide()
header(s, "An AI Operating System for Insurance", "We don't replace an insurer's software — we give it enterprise-grade AI through APIs.", section="Overview", kicker="The one-line idea")
two_col(s,
    [('H', 'Built today'),
     ('B', 'A multi-agent conversational advisory platform that educates and guides customers'),
     ('B', 'Personalised recommendation & scoring across curated insurance plans'),
     ('B', 'Real-time voice, memory, secure auth and a service API layer')],
    [('H', 'The vision'),
     ('B', 'Fraud, damage, OCR, identity & decision intelligence — delivered as APIs'),
     ('B', 'A modular “AI OS” insurers integrate, instead of rebuilding AI in-house'),
     ('B', 'Explainable, evidence-based, human-in-the-loop by design')])
caption(s, "Center hero diagram: an ‘AI OS’ core ringed by pluggable intelligence APIs feeding existing insurer systems.")
notes(s, N(
    "Frame the category: infrastructure/OS, integrated via APIs.",
    "Every insurer today is trying to build AI in-house and failing — it's expensive, "
    "opaque, and hard to trust. Aegis flips that: we become the AI layer they call. "
    "Today that layer is a live advisory and recommendation platform; the roadmap turns "
    "it into a full intelligence suite — fraud, damage, documents, identity, decisions.",
    [("Why an OS analogy?",
      "Like an OS, we sit between the insurer's applications and the underlying AI, "
      "offering standard, governed, explainable services any app can call.")]))

# 3 — TWO ALTITUDES
s = new_slide()
header(s, "One Project, Two Altitudes", "Engineered to satisfy an academic jury and an investment committee at once.", section="Overview", kicker="How to read this deck")
two_col(s,
    [('H', 'As a final-year engineering project'),
     ('B', 'Real multi-agent architecture, not a toy chatbot'),
     ('B', 'Grounded in peer-reviewed insurance-AI research'),
     ('B', 'Full-stack system: AI engine, backend, database, UI'),
     ('B', '200+ automated tests, CI, security hardening')],
    [('H', 'As a startup'),
     ('B', 'A platform + business model, not a single feature'),
     ('B', 'API-as-a-Service go-to-market for global insurers'),
     ('B', 'A 5-phase roadmap to an Insurance Foundation Model'),
     ('B', 'Clear separation of what is shipped vs. what is funded next')])
caption(s, "Split-screen: graduation cap ↔ rocket, sharing one core platform icon in the middle.")
notes(s, N(
    "Set expectations: this deck is honest about built vs. vision.",
    "I'll be precise throughout about what is running today versus what is on the roadmap "
    "— because in insurance, trust is the product. Everything I label ‘built’ you can see "
    "in the codebase; everything I label ‘vision’ is a funded next step, not a claim.",
    [("How much is actually working?",
      "The advisory, recommendation, voice, memory, auth and API layers are working and "
      "tested. Computer-vision fraud/damage/OCR are the roadmap — I'll mark them clearly.")]))

# =============================================================================
#  SECTION B — VISION / MISSION / GOALS
# =============================================================================

# 4 — VISION
s = new_slide()
header(s, "Vision", section="Direction", kicker="Where we are going")
_, tf = textbox(s, Inches(0.95), Inches(2.1), Inches(11.4), Inches(2.2))
p = para(tf, first=True, line=1.15, space_after=0)
run(p, "To build the world's most intelligent AI infrastructure platform that transforms "
       "insurance through ", 24, WHITE)
run(p, "explainable, secure, and trustworthy AI", 24, CYAN, bold=True)
run(p, " — preventing fraud, accelerating claims, improving customer experience, and letting "
       "every insurer integrate enterprise-grade AI through APIs.", 24, WHITE)
blocks(s, [('K', 'Prevent', 'Stop fraud before payout, with evidence.'),
           ('K', 'Accelerate', 'Turn multi-day claims into minutes.'),
           ('K', 'Trust', 'Every decision explainable and auditable.')],
       y=Inches(4.7), size=14)
notes(s, N("State the north star in one sentence.",
    "Our vision is simple to say and hard to build: make insurance AI you can actually "
    "trust — explainable, secure, and available to every insurer as infrastructure."))

# 5 — MISSION
s = new_slide()
header(s, "Mission", "A modular enterprise platform delivering insurance intelligence as scalable APIs.", section="Direction", kicker="What we do")
grid = ["Fraud Intelligence", "Damage Intelligence", "OCR Intelligence",
        "Recommendation Intelligence", "Claims Intelligence", "Identity Verification",
        "Policy Intelligence", "Decision Intelligence", "Explainable AI"]
gx, gy = Inches(0.95), Inches(2.15)
cw, ch, mx, my = Inches(3.75), Inches(1.15), Inches(0.18), Inches(0.22)
for i, name in enumerate(grid):
    col, row = i % 3, i // 3
    x = gx + col * (cw + mx)
    y = gy + row * (ch + my)
    pnl = panel(s, x, y, cw, ch)
    tf = pnl.text_frame
    tf.word_wrap = True
    pp = para(tf, first=True, space_after=0, align=PP_ALIGN.CENTER)
    run(pp, name, 15, WHITE, bold=True, font=FONT_SB)
    status = GREEN if name in ("Recommendation Intelligence", "Policy Intelligence", "Explainable AI") else AMBER
    lbl = "BUILT (advisory)" if status is GREEN else "VISION / API"
    p2 = para(tf, space_before=3, space_after=0, align=PP_ALIGN.CENTER)
    run(p2, lbl, 9, status, bold=True)
caption(s, "3×3 capability grid with green = built, amber = roadmap; each tile becomes an API endpoint.")
notes(s, N("Show the capability surface as a set of API-shaped modules.",
    "Our mission is to deliver these nine intelligences as clean APIs. Three are live in "
    "advisory form today; the rest are the roadmap. The point is the shape: modular, "
    "composable, callable.",
    [("Nine modules is a lot — realistic?",
      "We don't ship all nine at once. Each is an independent API with its own models and "
      "roadmap; we sequence them by customer pull and data availability.")]))

# 6 — MAIN GOAL
s = new_slide()
header(s, "Main Goal", section="Direction", kicker="The thesis")
_, tf = textbox(s, Inches(0.95), Inches(2.2), Inches(11.4), Inches(1.8))
p = para(tf, first=True, line=1.15)
run(p, "Build an ", 26, WHITE)
run(p, "AI Operating System for Insurance", 26, CYAN, bold=True)
run(p, " that insurers worldwide integrate through APIs — ", 26, WHITE)
run(p, "augmenting", 26, WHITE, bold=True)
run(p, " their existing software, not replacing it.", 26, WHITE)
two_col(s,
    [('H', 'Why augment, not replace'),
     ('B', 'Insurers have decades of core systems — rip-and-replace fails'),
     ('B', 'APIs mean weeks-to-integrate, not years-to-migrate'),
     ('B', 'We own the AI layer; they keep their systems of record')],
    [('H', 'What that unlocks'),
     ('B', 'One integration → many intelligences'),
     ('B', 'Central governance, monitoring and explainability'),
     ('B', 'A defensible platform position, not a point tool')],
    y=Inches(4.0), h=Inches(2.6))
notes(s, N("Crystallise the strategy: integrate, don't replace.",
    "The single most important strategic choice we make: we do not ask insurers to replace "
    "anything. We sit alongside their core systems and give them AI through APIs. That's how "
    "you actually get into a conservative, regulated industry."))

# 7 — GOALS (short vs long)
s = new_slide()
header(s, "Short-Term & Long-Term Goals", section="Direction", kicker="The path")
two_col(s,
    [('H', 'Short-term — the MVP'),
     ('B', 'Executive AI orchestration'),
     ('B', 'Customer & Admin portals, authentication'),
     ('B', 'Recommendation + Policy intelligence (advisory)'),
     ('B', 'REST API surface + Explainable reasoning trace'),
     ('B2', 'Roadmap MVP add-ons: Fraud · OCR · Damage detection')],
    [('H', 'Long-term — the platform'),
     ('B', 'Enterprise SaaS + multi-tenant platform'),
     ('B', 'Insurance AI Marketplace of intelligences'),
     ('B', 'Global enterprise APIs & developer ecosystem'),
     ('B', 'Private Insurance Foundation Model'),
     ('B', 'An industry-standard intelligence platform')])
caption(s, "Horizontal timeline arrow from ‘MVP’ to ‘Industry Standard’, milestones as nodes.")
notes(s, N("Contrast near-term deliverables with the platform end-state.",
    "Short term, we harden the MVP and add the first computer-vision modules. Long term, "
    "this becomes multi-tenant SaaS, a marketplace, and eventually a foundation model "
    "trained specifically for insurance."))

# =============================================================================
#  SECTION C — PROBLEM
# =============================================================================

# 8 — INDUSTRY TODAY
s = new_slide()
header(s, "The Insurance Industry Runs on Trust It Can't Verify", "Four structural failures compound into leakage, delay, and distrust.", section="Problem", kicker="Problem landscape")
cards = [("FRAUD LEAKAGE", "Fake, duplicate and staged claims drain a large share of every payout pool.", AMBER),
         ("SLOW CLAIMS", "Manual verification and investigation stretch settlements to days or weeks.", CYAN),
         ("OPAQUE AI", "Black-box models make decisions no adjuster can explain or defend.", ROYAL),
         ("BROKEN TRUST", "Customers distrust insurers; insurers distrust claims. Everyone loses.", GREEN)]
gx, gy, cw, ch, mx = Inches(0.95), Inches(2.15), Inches(5.75), Inches(1.85), Inches(0.35)
for i, (t, d, c) in enumerate(cards):
    x = gx + (i % 2) * (cw + mx)
    y = gy + (i // 2) * (ch + Inches(0.3))
    pnl = panel(s, x, y, cw, ch)
    rect(s, x, y, Inches(0.09), ch, c)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.18); tf.margin_right = Inches(0.2)
    pp = para(tf, first=True, space_after=4)
    run(pp, t, 15, c, bold=True, font=FONT_SB)
    p2 = para(tf, space_after=0, line=1.05)
    run(p2, d, 13, MUTE)
caption(s, "Four-quadrant ‘cost of the status quo’; industry-estimate figures cited, labelled as estimates.")
notes(s, N("Set up the problem at industry scale before the deep dives.",
    "Insurance is a trust business running on processes that can't verify trust. Widely "
    "cited industry estimates put fraud leakage in the hundreds of billions of dollars a "
    "year globally, while honest customers wait days for claims. I'll cite estimates as "
    "estimates — but the direction is not in dispute.",
    [("Your market numbers — where from?",
      "Public industry estimates (e.g., Coalition Against Insurance Fraud, Swiss Re/Deloitte "
      "reports). I present them as cited estimates, never as our own measured data.")]))

# 9 — PROBLEM: CLAIMS FRAUD
s = new_slide()
header(s, "Problem 1 — Claims Fraud Is Organised, Not Occasional", section="Problem", kicker="Deep dive · fraud")
two_col(s,
    [('H', 'The failure modes'),
     ('B', 'Fake & duplicate claims filed across insurers'),
     ('B', 'Collusion rings: hospital · garage · agent · customer'),
     ('B', 'Staged accidents and inflated repair estimates'),
     ('B', 'Identity fraud and policy misuse')],
    [('H', 'Why it persists'),
     ('B', 'Verification is manual, siloed and time-boxed'),
     ('B', 'No cross-insurer signal to catch duplicates'),
     ('B', 'Adjusters lack tools to see relationships between actors'),
     ('H', 'Business impact'),
     ('B', 'Direct leakage + higher premiums for honest customers')])
caption(s, "Fraud-ring graph: nodes for customer/hospital/garage/agent, edges = suspicious co-occurrence.")
notes(s, N("Establish fraud as a networked, systemic problem — motivating a graph approach later.",
    "Fraud isn't one bad actor; it's networks — a garage, an agent, a customer, sometimes a "
    "clinic, repeating patterns across claims. Humans can't see those networks at scale. "
    "This is exactly why our fraud roadmap centres on graph intelligence.",
    [("Isn't fraud detection already solved?",
      "Rules and basic anomaly scoring exist and are commoditised. What's missing is "
      "networked, explainable, evidence-backed detection — that's our differentiation.")]))

# 10 — PROBLEM: EVIDENCE MANIPULATION
s = new_slide()
header(s, "Problem 2 — The Evidence Itself Is Manipulable", section="Problem", kicker="Deep dive · evidence integrity")
blocks(s, [
    ('B', 'Fake or forged documents — bills, prescriptions, FIRs, RCs'),
    ('B', 'Edited images — cloned dents, added damage, tampered metadata'),
    ('B', 'Old-damage reuse — a prior accident re-submitted as new'),
    ('B', 'Deepfake risk — synthetic photos, video and voice in claims'),
    ('B', 'GPS / location spoofing — faking where an incident occurred'),
    ('S',),
    ('H', 'Technical root cause'),
    ('P', 'Claims are still adjudicated on human eyeballing of digital evidence that modern '
          'tools can forge faster than an adjuster can catch — with no automated provenance, '
          'tamper, or reuse checks in the loop.'),
])
caption(s, "Before/after image pair with forensic overlay (ELA / metadata / duplicate-match heatmap).")
notes(s, N("Show that the evidence layer is broken — motivates CV forensics + provenance.",
    "Even when a claim is honest-looking, the evidence can be fabricated. Edited images, "
    "recycled old damage, forged bills, and increasingly deepfakes. There is no automated "
    "forensic check today. That's a whole product category by itself."))

# 11 — PROBLEM: SPEED & CX
s = new_slide()
header(s, "Problem 3 — Manual Process Kills Speed and Experience", section="Problem", kicker="Deep dive · speed & CX")
two_col(s,
    [('H', 'Current workflow'),
     ('B', 'Intake → manual document review → field/desk investigation'),
     ('B', 'Back-and-forth with the customer for missing evidence'),
     ('B', 'Human adjudication → payout — days to weeks'),
     ('H', 'Business impact'),
     ('B', 'High operating cost per claim; adjuster burnout')],
    [('H', 'Customer-side pain'),
     ('B', 'Opaque status, jargon-heavy policies, poor guidance'),
     ('B', 'One-size recommendations that ignore real need & budget'),
     ('B', 'Distrust at the exact moment of vulnerability (a loss event)'),
     ('H', 'Underserved users'),
     ('B', 'First-time buyers, seniors, rural & multilingual customers')])
caption(s, "Swimlane: today's multi-day manual claim vs. an Aegis-assisted minutes-scale flow.")
notes(s, N("Tie operational cost to customer experience — and to our mission's underserved users.",
    "Manual process is slow and expensive for the insurer and miserable for the customer. "
    "And it fails hardest for the people insurance already underserves — first-time buyers, "
    "seniors, rural and multilingual users. That underserved audience is where our advisory "
    "product already delivers today."))

# 12 — PROBLEM: AI TRUST
s = new_slide()
header(s, "Problem 4 — Today's AI Can't Be Trusted With Claims", section="Problem", kicker="Deep dive · the AI trust gap")
two_col(s,
    [('H', 'The black-box problem'),
     ('B', 'Models output a score with no reason a human can audit'),
     ('B', 'No evidence trail linking decision → data → rule'),
     ('B', 'Regulators & adjusters cannot defend the outcome'),
     ('B', 'Bias and errors are invisible until they cause harm')],
    [('H', 'What insurance actually needs'),
     ('B', 'Explainable, evidence-based decisions'),
     ('B', 'Human-in-the-loop on every high-stakes call'),
     ('B', 'Full audit logs and model monitoring'),
     ('B', 'Governance, security and compliance as first-class')])
caption(s, "Contrast panel: ‘black box → score’ vs. ‘Aegis → decision + evidence + explanation + audit’.")
notes(s, N("Name the gap our platform is designed around: trustworthy, explainable AI.",
    "Here's the deepest problem. Even accurate AI is useless in claims if it can't explain "
    "itself. Insurance is regulated and adversarial — every automated decision must be "
    "explainable, evidence-based, auditable, and human-supervised. Almost nobody builds AI "
    "that way. We do, from the ground up."))

# 13 — PROBLEM SUMMARY TABLE
s = new_slide()
header(s, "The Problem, Mapped", "Every failure below has a corresponding Aegis intelligence.", section="Problem", kicker="Summary")
rows = [
    ("Problem cluster", "Who it hurts", "Aegis answer"),
    ("Fake / duplicate / ring fraud", "Insurer payout pool", "Fraud + Network Intelligence"),
    ("Forged docs / edited & old images", "Claims integrity", "OCR + Damage + Evidence Intelligence"),
    ("Identity fraud / deepfake / spoofing", "Underwriting & claims", "Identity + Liveness Intelligence"),
    ("Slow, manual, costly claims", "Insurer + customer", "Claims + Investigation Agents"),
    ("Black-box, untrusted AI", "Regulator + adjuster", "Explainable + Decision Intelligence"),
    ("Poor guidance & recommendations", "Underserved customers", "Advisory + Recommendation (LIVE)"),
]
# render table manually
from_x, from_y = Inches(0.95), Inches(2.05)
cw = [Inches(4.4), Inches(3.1), Inches(4.1)]
rh = Inches(0.62)
for r, row in enumerate(rows):
    x = from_x
    for c, cell in enumerate(row):
        head = (r == 0)
        fill = NAVY3 if head else (NAVY2 if r % 2 else NAVY)
        cellsp = rect(s, x, from_y + r * rh, cw[c], rh, fill, line=LINE)
        tf = cellsp.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.15); tf.margin_top = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = para(tf, first=True, space_after=0)
        live = "(LIVE)" in cell
        run(p, cell.replace(" (LIVE)", ""), 12.5 if head else 12,
            CYAN if head else (GREEN if live else WHITE), bold=head or (c == 2))
        if live:
            run(p, "  ● LIVE", 10, GREEN, bold=True)
        x += cw[c]
notes(s, N("One glance: problems on the left, our modules on the right; note recommendation is already live.",
    "This is the whole thesis on one slide. Left column, the problems. Right column, the "
    "Aegis intelligence that answers each. Notice the last row is already live — advisory "
    "and recommendation — and everything else is a sequenced roadmap off the same platform."))

# =============================================================================
#  SECTION D — RESEARCH PAPER CHAPTER
# =============================================================================

# 14 — RESEARCH CHAPTER INTRO
s = new_slide()
header(s, "Research Foundation", section="Research", kicker="Chapter · bridging research and reality")
panel(s, Inches(0.95), Inches(2.1), Inches(11.4), Inches(1.4))
_, tf = textbox(s, Inches(1.2), Inches(2.3), Inches(11.0), Inches(1.1), anchor=MSO_ANCHOR.MIDDLE)
p = para(tf, first=True, space_after=0)
run(p, "“AI Revolution in Insurance: Bridging Research and Reality”", 20, WHITE, bold=True, font=FONT_SB)
p2 = para(tf, space_before=4, space_after=0)
run(p2, "Anchor paper for our gap analysis — what the literature promises vs. what industry ships.", 13, MUTE)
blocks(s, [('H', 'This chapter covers'),
           ('B', 'Paper overview, objectives and key findings'),
           ('B', 'Current industry status & already-existing technologies'),
           ('B', 'Research gaps, challenges and limitations'),
           ('B', 'How Aegis AI converts each gap into a product decision')],
       y=Inches(3.9))
notes(s, N("Signal academic rigour: our product is a response to a specific research gap.",
    "We didn't start from ‘let's build AI'. We started from the literature. This paper — "
    "Bridging Research and Reality — maps what research promises against what insurers "
    "actually deploy. The delta between the two is our product."))

# 15 — RESEARCH OBJECTIVES & FINDINGS
s = new_slide()
header(s, "What the Research Set Out to Answer — and Found", section="Research", kicker="Objectives & findings")
two_col(s,
    [('H', 'Research objectives'),
     ('B', 'Map AI adoption across the insurance value chain'),
     ('B', 'Assess maturity of fraud, claims and CX automation'),
     ('B', 'Identify blockers to real-world deployment'),
     ('B', 'Surface open research directions')],
    [('H', 'Key findings'),
     ('B', 'AI works in pilots; stalls in production'),
     ('B', 'Trust, explainability & data quality are the real blockers'),
     ('B', 'Point solutions exist; integrated platforms don’t'),
     ('B', 'Regulation demands transparency few systems provide')])
caption(s, "Value-chain heatmap: AI maturity per stage (quote/underwrite/service/claim/fraud).")
notes(s, N("Summarise objectives and findings crisply.",
    "The findings are striking: AI demos beautifully and dies in production. The blockers "
    "aren't model accuracy — they're trust, explainability, data quality and integration. "
    "That reframes the opportunity from ‘better models' to ‘deployable, trustworthy platform'."))

# 16 — WHAT ALREADY EXISTS
s = new_slide()
header(s, "What Already Exists — and Is Commoditised", "Being honest about prior art sharpens our differentiation.", section="Research", kicker="Existing technologies")
grid = [("Chatbots", "Conversational FAQ & sales bots"),
        ("OCR", "Document text extraction"),
        ("Fraud scoring", "Rules + basic anomaly models"),
        ("Recommendation", "Product suggestion engines"),
        ("Claims automation", "Workflow / RPA pipelines")]
gx, gy, cw, ch, mx = Inches(0.95), Inches(2.3), Inches(2.24), Inches(2.0), Inches(0.13)
for i, (t, d) in enumerate(grid):
    x = gx + i * (cw + mx)
    pnl = panel(s, x, gy, cw, ch)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.15); tf.margin_right = Inches(0.15); tf.margin_top = Inches(0.2)
    p = para(tf, first=True, space_after=6, align=PP_ALIGN.CENTER)
    run(p, t, 14, WHITE, bold=True, font=FONT_SB)
    p2 = para(tf, space_after=8, align=PP_ALIGN.CENTER, line=1.05)
    run(p2, d, 11, MUTE)
    p3 = para(tf, space_after=0, align=PP_ALIGN.CENTER)
    run(p3, "COMMODITY", 9, AMBER, bold=True)
_, tf2 = textbox(s, Inches(0.95), Inches(4.7), Inches(11.4), Inches(1.2))
p = para(tf2, first=True, line=1.15)
run(p, "None of these decide, explain, or investigate. They are features. ", 15, WHITE)
run(p, "Aegis composes them into governed, explainable decisions.", 15, CYAN, bold=True)
notes(s, N("Disarm the ‘this already exists' objection by naming prior art first.",
    "Let me be the first to say it: chatbots, OCR, fraud scoring, recommendation and claims "
    "automation all exist and are commoditised. If that's all we were, we'd be late. But "
    "these are features. None of them make an explainable, evidence-based decision or run an "
    "investigation. That composition is the whitespace.",
    [("How are you different from existing insurtech?",
      "They sell point features. We sell the decision + evidence + explanation layer that "
      "orchestrates those features — plus the API platform to deliver it.")]))

# 17 — GAPS / CHALLENGES / LIMITATIONS
s = new_slide()
header(s, "Research Gaps, Challenges & Limitations", section="Research", kicker="The whitespace")
blocks(s, [
    ('K', 'Explainability gap', 'Scores without reasons no regulator or adjuster can defend.'),
    ('K', 'Decision gap', 'Detection exists; end-to-end evidence-based decisioning does not.'),
    ('K', 'Integration gap', 'Point tools; no unified, API-first AI platform for insurers.'),
    ('K', 'Trust gap', 'No human-in-the-loop, audit and governance baked in.'),
    ('K', 'Data gap', 'Fragmented, low-quality, non-shared claims data.'),
    ('K', 'Investigation gap', 'No autonomous evidence collection & network analysis.'),
], size=15, gap=11)
caption(s, "Six-gap ‘bridge’ graphic — research on one bank, reality on the other, Aegis as the bridge.")
notes(s, N("List the exact gaps we will map to features next.",
    "Six gaps: explainability, decision, integration, trust, data and investigation. Keep "
    "these six in mind — the next slide maps each one directly to an Aegis capability. That "
    "one-to-one mapping is what makes this research-driven, not feature-driven."))

# 18 — HOW AEGIS DIFFERS
s = new_slide()
header(s, "How Aegis AI Is Different", "We don't add another model — we add the missing layer.", section="Research", kicker="Differentiation")
diff = ["Executive AI", "Decision Intelligence", "Explainable AI", "Investigation AI",
        "Evidence Intelligence", "Enterprise API Platform", "AI Infrastructure", "Human-in-the-Loop", "Trustworthy AI"]
gx, gy, cw, ch, mx, my = Inches(0.95), Inches(2.2), Inches(3.75), Inches(1.05), Inches(0.18), Inches(0.2)
for i, name in enumerate(diff):
    col, row = i % 3, i // 3
    x = gx + col * (cw + mx); y = gy + row * (ch + my)
    pnl = panel(s, x, y, cw, ch, rgb=NAVY3)
    rect(s, x, y, cw, Inches(0.07), CYAN)
    tf = pnl.text_frame; tf.word_wrap = True
    p = para(tf, first=True, space_after=0, align=PP_ALIGN.CENTER)
    run(p, name, 14.5, WHITE, bold=True, font=FONT_SB)
notes(s, N("Present the nine differentiators as the missing ‘decision & trust' layer.",
    "Our nine differentiators aren't more detectors — they're the layer above detection: an "
    "Executive AI that orchestrates, decision and evidence intelligence, explainability, "
    "investigation, and the enterprise API to ship it — all human-in-the-loop. That's the "
    "bridge the paper says is missing."))

# 19 — GAP → PRODUCT MAPPING
s = new_slide()
header(s, "From Research Gap to Product Decision", section="Research", kicker="Gap → answer")
rows = [
    ("Research gap", "Aegis product decision", "Status"),
    ("Explainability", "Reasoning trace on every answer (ThinkingEngine)", "LIVE (advisory)"),
    ("Decision", "Decision Intelligence agent + rules-driven KB", "Roadmap"),
    ("Integration", "API-first, provider-agnostic architecture", "LIVE (core)"),
    ("Trust", "RBAC, tenant isolation, audit, human-in-the-loop", "LIVE / Roadmap"),
    ("Investigation", "Evidence + Investigation agents", "Roadmap"),
    ("Data", "Insurance Knowledge Graph + Foundation Model", "Vision"),
]
from_x, from_y = Inches(0.95), Inches(2.05); cw = [Inches(3.0), Inches(5.9), Inches(2.6)]; rh = Inches(0.62)
for r, row in enumerate(rows):
    x = from_x
    for c, cell in enumerate(row):
        head = (r == 0)
        fill = NAVY3 if head else (NAVY2 if r % 2 else NAVY)
        cellsp = rect(s, x, from_y + r * rh, cw[c], rh, fill, line=LINE)
        tf = cellsp.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.15); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = para(tf, first=True, space_after=0)
        col = CYAN if head else WHITE
        if c == 2 and not head:
            col = GREEN if cell.startswith("LIVE") else (AMBER if cell == "Roadmap" else ROYAL)
        run(p, cell, 12.5 if head else 12, col, bold=head or c == 2)
        x += cw[c]
notes(s, N("Prove the research-to-build traceability, honestly labelled.",
    "Every gap becomes a concrete decision, and I mark exactly what's live versus roadmap "
    "versus vision. Explainability and integration are live in the core today; decision and "
    "investigation are the funded next steps; the knowledge graph and foundation model are "
    "the long-horizon vision."))

# =============================================================================
#  SECTION E — SOLUTION: BUILT TODAY
# =============================================================================

# 20 — SOLUTION OVERVIEW
s = new_slide()
header(s, "Aegis AI — The Platform", "Five layers, one governed AI operating system.", section="Solution", kicker="Solution overview")
layers = [("Experience", "Customer · Admin · (Enterprise) portals, voice & chat", GREEN),
          ("Orchestration", "Executive AI routes intent to the right specialist agent", GREEN),
          ("Intelligence", "Advisory & Recommendation today → Fraud/Damage/OCR next", AMBER),
          ("Memory & Data", "Profiles, conversation & recommendation memory", GREEN),
          ("Platform", "Auth, RBAC, tenant isolation, API layer, provider-agnostic LLM", GREEN)]
gy = Inches(2.15)
for i, (t, d, c) in enumerate(layers):
    y = gy + i * Inches(0.92)
    pnl = panel(s, Inches(0.95), y, Inches(11.4), Inches(0.78))
    rect(s, Inches(0.95), y, Inches(0.12), Inches(0.78), c)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.3); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = para(tf, first=True, space_after=0)
    run(p, f"{t}   ", 15, c, bold=True, font=FONT_SB)
    run(p, d, 13, MUTE)
notes(s, N("Give the mental model: a 5-layer stack, mostly built.",
    "Think of Aegis as five layers. Experience, orchestration, intelligence, memory, and "
    "platform. Four of the five are live today; the intelligence layer is where the "
    "computer-vision roadmap plugs in. The architecture already assumes those modules."))

# 21 — BUILT TODAY
s = new_slide()
header(s, "What Is Actually Built Today", "Running, tested code — not slideware.", section="Solution", kicker="● Built today", )
two_col(s,
    [('H', 'AI engine (Python · FastAPI)'),
     ('B', 'Multi-agent system: Executive + 4 domain advisors'),
     ('B', 'Central Orchestrator, intent routing, safe transfers'),
     ('B', 'Layer-3 memory: profile · conversation · recommendation'),
     ('B', 'Recommendation & scoring across 36 curated plans'),
     ('B', 'Provider-agnostic LLM layer (Ollama / Gemini / OpenAI)')],
    [('H', 'Product (Node · Next.js)'),
     ('B', 'Secure backend: JWT httpOnly, RBAC, tenant isolation'),
     ('B', 'Real-time voice (SSE streaming) + live reasoning trace'),
     ('B', 'Consumer & Admin dashboards, leads, policies'),
     ('B', 'Service API layer + Socket.io realtime'),
     ('B', '200+ automated tests, CI, security hardening')])
caption(s, "Screenshot strip: advisor chat, recommendation cards, admin dashboard, voice UI.")
notes(s, N("This is the credibility slide — everything here is verifiable in the repo.",
    "This is the slide I most want you to hold me to. Every line here is in the codebase and "
    "under test. A working multi-agent AI engine, a secure full-stack product, voice, "
    "memory, recommendations, dashboards, and 200-plus tests. This is a real system, today.",
    [("Can you demo it live?",
      "Yes — the advisor chat, recommendation flow, voice and dashboards all run locally. "
      "I can show the multi-agent transfer and the live reasoning trace on request.")]))

# 22 — SYSTEM ARCHITECTURE (drawn)
s = new_slide()
header(s, "System Architecture", "Request flows through one router; agents stay isolated; memory is the only cross-domain path.", section="Solution", kicker="Architecture")


def abox(x, y, w, h, title, sub, c=NAVY3, tc=WHITE):
    pnl = panel(s, x, y, w, h, rgb=c)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = para(tf, first=True, space_after=0, align=PP_ALIGN.CENTER)
    run(p, title, 12.5, tc, bold=True, font=FONT_SB)
    if sub:
        p2 = para(tf, space_before=1, space_after=0, align=PP_ALIGN.CENTER)
        run(p2, sub, 9.5, MUTE)
    return pnl


def arrow(x, y, w=Inches(0.5)):
    sp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.28))
    solid(sp, ROYAL)


y0 = Inches(2.5)
abox(Inches(0.95), y0, Inches(1.9), Inches(0.9), "Frontend", "Next.js UI · Voice")
arrow(Inches(2.95), y0 + Inches(0.31))
abox(Inches(3.55), y0, Inches(1.9), Inches(0.9), "Backend API", "Node · Express · Prisma")
arrow(Inches(5.55), y0 + Inches(0.31))
abox(Inches(6.15), y0, Inches(2.2), Inches(0.9), "Central Orchestrator", "Intent · Interrupt · Transfer", c=RGBColor(0x1B, 0x2A, 0x55))
# agents row
agents = ["Executive", "Health", "Motor", "Travel", "Property"]
ax, ay, aw = Inches(0.95), Inches(4.0), Inches(1.55)
for i, a in enumerate(agents):
    abox(ax + i * (aw + Inches(0.14)), ay, aw, Inches(0.8), a, "Agent", c=NAVY2)
# down arrow from orchestrator to agents
dn = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(7.1), y0 + Inches(0.95), Inches(0.28), Inches(0.5))
solid(dn, ROYAL)
# memory + llm + kb
abox(Inches(9.0), y0, Inches(3.35), Inches(0.9), "Memory Orchestrator", "Profile · Conversation · Rec", c=RGBColor(0x1B, 0x2A, 0x55))
abox(Inches(0.95), Inches(5.2), Inches(5.5), Inches(0.8), "Rule-Driven Knowledge Base (Layers 1–5)", "Domain · routing · scoring · risk · executive", c=NAVY2)
abox(Inches(6.7), Inches(5.2), Inches(2.5), Inches(0.8), "LLM Layer", "Ollama / Gemini / OpenAI", c=NAVY2)
abox(Inches(9.4), Inches(5.2), Inches(2.95), Inches(0.8), "Roadmap: CV / Fraud / OCR APIs", "pluggable intelligence", c=RGBColor(0x33, 0x2A, 0x12), tc=AMBER)
notes(s, N("Walk the request path and stress agent isolation + single router.",
    "A request comes from the UI, through the secure backend, into the Central Orchestrator "
    "— the only router. It detects intent, handles interruptions, and dispatches to one "
    "isolated specialist agent. Agents never read each other's memory except through the "
    "Memory Orchestrator. The knowledge base and LLM layer sit underneath. And on the right, "
    "the roadmap CV/fraud/OCR modules plug into the same orchestration — the architecture is "
    "already built to receive them.",
    [("What stops one agent leaking another's data?",
      "Agent isolation is enforced: each has its own memory namespace, and the only "
      "cross-domain path is a sanctioned export in the Memory Orchestrator.")]))

# 23 — ORCHESTRATION DEEP DIVE
s = new_slide()
header(s, "Multi-Agent Orchestration", "The crown jewel: consented, auditable hand-offs between specialists.", section="Solution", kicker="Deep dive · orchestration")
two_col(s,
    [('H', 'How it works'),
     ('B', 'Executive AI detects intent and picks the specialist'),
     ('B', 'Interrupt detector catches mid-flow topic switches'),
     ('B', 'Transfers are user-consented, never silent'),
     ('B', 'Each agent has its own config, tools and memory namespace')],
    [('H', 'Why it matters'),
     ('B', 'Mirrors a real brokerage: a lead who routes to experts'),
     ('B', 'Isolation contains risk & keeps domains explainable'),
     ('B', 'The same router will dispatch fraud/damage/OCR agents'),
     ('B', 'Governance & audit live at the orchestration layer')])
caption(s, "Sequence diagram: user → Executive → suggest_transfer → consent → specialist → memory sync.")
notes(s, N("Explain the consented-transfer model and why it generalises to future agents.",
    "This is what makes us a platform, not a bot. The Executive AI behaves like a smart front "
    "desk: it understands intent and routes you to the right specialist, but only with your "
    "consent, and it can handle you changing your mind mid-conversation. Crucially, the exact "
    "same mechanism will route to a fraud agent or a damage agent tomorrow."))

# 24 — AGENTS ROSTER (drawn grid)
s = new_slide()
header(s, "The AI Agents", "Five live specialists today; six investigation agents on the roadmap.", section="Solution", kicker="Agent roster")
live = [("Executive AI", "Routing & oversight"), ("Health Advisor", "Health cover"),
        ("Motor Advisor", "Vehicle cover"), ("Travel Advisor", "Trip cover"),
        ("Property Advisor", "Home cover")]
future = [("Fraud Agent", "Anomaly + network"), ("OCR Agent", "Docs & forgery"),
          ("Damage Agent", "Vision estimate"), ("Claim Agent", "Adjudication"),
          ("Evidence Agent", "Provenance"), ("Investigation Agent", "Autonomous probe")]
_, tfl = textbox(s, Inches(0.95), Inches(1.95), Inches(5), Inches(0.4))
run(para(tfl, first=True, space_after=0), "● LIVE TODAY", 12, GREEN, bold=True)
for i, (t, d) in enumerate(live):
    y = Inches(2.4) + i * Inches(0.82)
    pnl = panel(s, Inches(0.95), y, Inches(5.4), Inches(0.7), rgb=NAVY3)
    rect(s, Inches(0.95), y, Inches(0.09), Inches(0.7), GREEN)
    tf = pnl.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.25)
    p = para(tf, first=True, space_after=0)
    run(p, t + "   ", 14, WHITE, bold=True, font=FONT_SB); run(p, d, 12, MUTE)
_, tfr = textbox(s, Inches(6.95), Inches(1.95), Inches(5), Inches(0.4))
run(para(tfr, first=True, space_after=0), "○ ROADMAP", 12, AMBER, bold=True)
for i, (t, d) in enumerate(future):
    y = Inches(2.4) + i * Inches(0.68)
    pnl = panel(s, Inches(6.95), y, Inches(5.4), Inches(0.58), rgb=NAVY2)
    rect(s, Inches(6.95), y, Inches(0.09), Inches(0.58), AMBER)
    tf = pnl.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.25)
    p = para(tf, first=True, space_after=0)
    run(p, t + "   ", 13, WHITE, bold=True); run(p, d, 11, MUTE)
notes(s, N("Show a real, growing agent org-chart; be explicit about live vs roadmap.",
    "Five specialist agents run today under the Executive AI. The roadmap adds an "
    "investigation cluster — fraud, OCR, damage, claim, evidence and investigation agents — "
    "onto the identical orchestration fabric. We're not redesigning; we're extending."))

# 25 — MEMORY
s = new_slide()
header(s, "Memory & Personalisation", "The system remembers the customer — safely and per-domain.", section="Solution", kicker="Deep dive · memory")
two_col(s,
    [('H', 'Layer-3 memory engine'),
     ('B', 'Customer profile persistence across sessions'),
     ('B', 'Per-domain conversation memory namespaces'),
     ('B', 'Recommendation context cached on a profile hash'),
     ('B', 'Cross-domain sync only via Memory Orchestrator')],
    [('H', 'Why it matters'),
     ('B', 'Continuity: no repeating yourself to each advisor'),
     ('B', 'Personalisation grounded in real, stored context'),
     ('B', 'Isolation keeps domains private and auditable'),
     ('B', 'Foundation for future learning & knowledge graph')])
caption(s, "Memory namespace diagram: {domain}_{customer_id} stores syncing through one orchestrator.")
notes(s, N("Position memory as both a UX win and a governance boundary.",
    "Memory does two jobs: it makes the experience feel personal and continuous, and it's a "
    "security boundary. Each domain has its own namespace; the only way data crosses domains "
    "is a sanctioned, audited export. That's how you personalise without leaking."))

# 26 — RECOMMENDATION
s = new_slide()
header(s, "Recommendation & Scoring Intelligence", "Budget-first, needs-aware guidance — live today.", section="Solution", kicker="● Built today")
two_col(s,
    [('H', 'What it does'),
     ('B', '36 curated plans across 4 categories, independently reviewed'),
     ('B', 'Scoring, risk and comparison engines match need → plan'),
     ('B', 'Confirmation-gated, explainable suggestions'),
     ('B', 'Multi-plan comparison surfaced as clear cards')],
    [('H', 'Mission alignment'),
     ('B', 'Plain-language, budget-first framing'),
     ('B', 'No jargon assumed; designed for first-time buyers'),
     ('B', 'Only verifiable facts shown (trust-integrity guardrails)'),
     ('B', 'Multilingual, low-bandwidth, senior-friendly UX')])
caption(s, "Recommendation flow: profile → score → risk → top plans with reason-for-fit chips.")
notes(s, N("Tie the recommendation engine to the mission and to trust integrity.",
    "Our recommendation engine is live and, importantly, honest — it only shows facts we can "
    "verify from the catalogue, with visible guardrails against invented claims. It's "
    "budget-first and plain-language, built for the people insurance usually underserves."))

# 27 — VOICE
s = new_slide()
header(s, "Voice & Real-Time Streaming", "A voice-first, low-bandwidth channel for rural and senior users.", section="Solution", kicker="● Built today")
two_col(s,
    [('H', 'Capabilities'),
     ('B', 'SSE token streaming — users see progress instantly'),
     ('B', 'Voice input & spoken responses (useVoice)'),
     ('B', 'Live reasoning trace as the agent thinks'),
     ('B', 'Graceful reduced-motion & accessibility support')],
    [('H', 'Why voice matters here'),
     ('B', 'Reaches non-readers and low-literacy users'),
     ('B', 'Works on low bandwidth and modest devices'),
     ('B', 'Natural for Tamil / Thanglish / English code-switching'),
     ('B', 'Turns insurance from a form into a conversation')])
caption(s, "Voice waveform + streaming transcript + a ‘thinking steps' panel.")
notes(s, N("Voice is an accessibility and reach story, not a gimmick.",
    "Voice and streaming aren't decoration — they're how we reach rural, senior and "
    "multilingual customers who won't fill a web form. The system streams its answer and even "
    "shows its reasoning live, which builds trust in the moment."))

# 28 — EXPLAINABILITY
s = new_slide()
header(s, "Explainability — Today and Next", "Trust starts with showing the work.", section="Solution", kicker="Explainable AI")
two_col(s,
    [('H', 'Live today'),
     ('B', 'ThinkingEngine streams the agent’s reasoning steps'),
     ('B', 'Recommendations show reason-for-fit, not just a pick'),
     ('B', 'Deterministic, rule-driven knowledge base underneath')],
    [('H', 'Roadmap'),
     ('B', 'Evidence-linked decisions (data → rule → outcome)'),
     ('B', 'Model explainability (SHAP / attention) for ML modules'),
     ('B', 'Human-in-the-loop review + full decision audit trail')])
caption(s, "‘Decision receipt' mock: verdict + the evidence and rules that produced it.")
notes(s, N("Show explainability is a live principle, deepening into evidence-based decisions.",
    "Explainability isn't a future promise for us — it's live in a basic form: the agent "
    "shows its reasoning, and recommendations explain themselves. The roadmap deepens this "
    "into evidence-linked decision receipts and model-level explanations for the ML modules."))

# 29 — SECURITY
s = new_slide()
header(s, "Trust, Security & Governance", "Built for a regulated, adversarial industry from day one.", section="Solution", kicker="● Built today")
blocks(s, [
    ('K', 'Auth', 'JWT in httpOnly cookies; Bearer for API clients; public registration is always customer.'),
    ('K', 'RBAC', 'Role-restricted routes; least-privilege access.'),
    ('K', 'Tenant isolation', 'Every user-data query scoped by user; cross-tenant leakage treated as critical.'),
    ('K', 'Engine gating', 'AI engine behind a constant-time internal API key; fail-closed in production.'),
    ('K', 'Input safety', 'Bounded, validated inputs (Node schemas + Pydantic); generic production errors.'),
    ('K', 'Roadmap', 'Dedicated security, compliance & AI-governance layers; model monitoring.'),
], size=14.5, gap=10)
notes(s, N("Security is a feature in insurance — lead with it.",
    "In insurance, security and governance are product features, not afterthoughts. Auth, "
    "RBAC, tenant isolation, a gated AI engine, bounded inputs — all live. The roadmap adds "
    "formal compliance and AI-governance layers as we approach enterprise deployment."))

# 30 — DASHBOARDS
s = new_slide()
header(s, "Portals & Dashboards", section="Solution", kicker="Experience layer")
cards = [("Customer Portal", "Advisory chat, voice, recommendations, policies", GREEN, "LIVE"),
         ("Admin Portal", "Operations, users, monitoring views", GREEN, "LIVE"),
         ("Enterprise Console", "Multi-tenant insurer control plane", AMBER, "ROADMAP"),
         ("Developer Portal", "API keys, docs, SDKs, usage & billing", AMBER, "ROADMAP")]
gx, gy, cw, ch, mx, my = Inches(0.95), Inches(2.15), Inches(5.6), Inches(1.9), Inches(0.2), Inches(0.25)
for i, (t, d, c, tag) in enumerate(cards):
    x = gx + (i % 2) * (cw + mx); y = gy + (i // 2) * (ch + my)
    pnl = panel(s, x, y, cw, ch)
    rect(s, x, y, Inches(0.09), ch, c)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.2)
    p = para(tf, first=True, space_after=4)
    run(p, t, 16, WHITE, bold=True, font=FONT_SB)
    run(p, f"   {tag}", 10, c, bold=True)
    p2 = para(tf, space_after=0, line=1.05); run(p2, d, 13, MUTE)
notes(s, N("Two portals live, two are the enterprise roadmap.",
    "Customer and admin portals are live. The enterprise console and developer portal — the "
    "surfaces that make this multi-tenant SaaS — are the roadmap, and they're what a design "
    "partner would help us shape."))

# =============================================================================
#  SECTION F — VISION / ROADMAP FEATURES
# =============================================================================

# 31 — VISION MATRIX (current vs vision, drawn)
s = new_slide()
header(s, "From Advisory Platform to Insurance AI OS", "The honest map of what's shipped and what's funded next.", section="Vision", kicker="Built vs. vision")
# two columns
_, tfl = textbox(s, Inches(1.1), Inches(2.0), Inches(5), Inches(0.4))
run(para(tfl, first=True, space_after=0), "● BUILT TODAY", 13, GREEN, bold=True)
_, tfr = textbox(s, Inches(7.1), Inches(2.0), Inches(5), Inches(0.4))
run(para(tfr, first=True, space_after=0), "○ VISION / ROADMAP", 13, AMBER, bold=True)
builtl = ["Multi-agent advisory", "Recommendation & scoring", "Voice + streaming",
          "Memory & personalisation", "Auth · RBAC · isolation", "Service API + reasoning trace"]
visionl = ["Fraud & Network Intelligence", "Damage (CV) & 3D estimate", "OCR & document forgery",
           "Identity · liveness · deepfake", "Evidence & Investigation agents", "Foundation model + knowledge graph"]
for i, t in enumerate(builtl):
    y = Inches(2.5) + i * Inches(0.62)
    pnl = panel(s, Inches(1.1), y, Inches(5.2), Inches(0.5), rgb=NAVY3)
    rect(s, Inches(1.1), y, Inches(0.08), Inches(0.5), GREEN)
    tf = pnl.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.22)
    run(para(tf, first=True, space_after=0), t, 13, WHITE, bold=True)
for i, t in enumerate(visionl):
    y = Inches(2.5) + i * Inches(0.62)
    pnl = panel(s, Inches(7.1), y, Inches(5.2), Inches(0.5), rgb=NAVY2)
    rect(s, Inches(7.1), y, Inches(0.08), Inches(0.5), AMBER)
    tf = pnl.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.22)
    run(para(tf, first=True, space_after=0), t, 13, WHITE, bold=True)
notes(s, N("The reference slide for honesty — return here if a jury pushes on scope.",
    "If you remember one slide, make it this one. Left is shipped and tested. Right is the "
    "funded roadmap. The reason I can promise the right column is that the left column's "
    "architecture — orchestration, memory, APIs, governance — was built to carry it."))


# vision feature template
def feature_slide(idx, title, kicker, problem, solution, models, workflow, benefit, revenue, extra_notes=""):
    s = new_slide()
    header(s, title, section="Vision", kicker=kicker)
    pill(s, Inches(11.0), Inches(0.62), "VISION / ROADMAP", AMBER, RGBColor(0x33, 0x2A, 0x12), w=Inches(1.9))
    two_col(s,
        [('H', 'Problem'), ('P', problem),
         ('H', 'Our solution'), ('P', solution),
         ('H', 'Enterprise benefit'), ('P', benefit)],
        [('H', 'ML / DL / CV models'), *[('B2', m) for m in models],
         ('H', 'Workflow'), ('P', workflow),
         ('H', 'Revenue'), ('P', revenue)])
    notes(s, N(f"Sell {title} as a self-contained API business with clear models and value.",
        f"{problem} {solution} Commercially, {revenue.lower()}",
        [("Do you have the data/models for this yet?",
          "Not deployed yet — this is roadmap. The architecture, agent fabric and API layer "
          "that host it are built; the models are the funded build. " + extra_notes)]))
    return s


# 32 — Damage Intelligence
feature_slide(32, "Vision — Damage Intelligence", "Computer vision · claims",
    "Vehicle & property damage is assessed by human eyeballing of photos — slow, inconsistent, and gameable with edited or old images.",
    "A CV pipeline that detects, segments and estimates damage severity from photos/video, flags reused or edited images, and produces an explainable repair estimate.",
    ["YOLO — damage part/region detection", "SAM — damage segmentation masks",
     "CNN / Vision Transformer — severity grading", "3D reconstruction — dent depth & extent",
     "Image-forensics net — edit/duplicate detection"],
    "Upload → detect parts → segment damage → grade severity → forensic check → explainable estimate → human review.",
    "Faster, consistent, fraud-resistant estimates; lower loss-adjustment expense.",
    "Per-inspection API pricing + enterprise seats for claims teams.")

# 33 — Fraud Intelligence
feature_slide(33, "Vision — Fraud & Network Intelligence", "ML + graph · fraud",
    "Fraud is networked (garage–agent–customer rings) and adaptive; rules and single-claim scoring miss coordinated abuse.",
    "A hybrid engine: anomaly detection on each claim plus a Graph Neural Network over the actor network to surface rings, with an explained risk score and evidence.",
    ["Isolation Forest / Autoencoder — anomaly", "XGBoost / LightGBM — supervised risk",
     "Graph Neural Network — ring detection", "Behaviour & Network Intelligence features",
     "SHAP — per-decision explanation"],
    "Claim + entities → features + graph → anomaly & GNN scores → explained risk + evidence → investigator review.",
    "Reduced leakage, higher recovery, defensible decisions for regulators.",
    "Usage-based scoring API + premium ‘network intelligence' tier.")

# 34 — OCR & Forgery
feature_slide(34, "Vision — OCR & Document Forgery Intelligence", "Document AI",
    "Claims run on documents — bills, prescriptions, RCs, FIRs — that are trivially forged and only spot-checked by humans.",
    "Intelligent OCR that extracts and structures fields, cross-checks them for consistency, and runs forgery/tamper forensics on the document image itself.",
    ["Vision Transformer / CNN-OCR — extraction", "Layout models — structure & fields",
     "Forgery-detection net — tamper/clone", "LLM — cross-field consistency reasoning"],
    "Doc → OCR + layout → structured fields → forgery forensics → consistency check → flag + confidence.",
    "Straight-through processing for clean docs; instant flags on suspicious ones.",
    "Per-page OCR pricing + forgery-check add-on.")

# 35 — Identity & Liveness
feature_slide(35, "Vision — Identity, Liveness & Deepfake Defence", "Biometric integrity",
    "Identity fraud, deepfakes and GPS spoofing let bad actors impersonate customers and fake where/when an incident happened.",
    "Live camera & video verification with liveness detection, face matching, deepfake detection, and location/metadata anti-spoofing at claim intake.",
    ["Liveness detection — anti-spoof", "Face-match embeddings — identity",
     "Deepfake-detection net — synthetic media", "Metadata / GPS anti-spoof checks"],
    "Capture → liveness → identity match → deepfake & location checks → verified / challenge.",
    "Stops impersonation at the door; strengthens KYC and claim intake.",
    "Per-verification API + enterprise KYC bundle.")

# 36 — Evidence & Investigation
feature_slide(36, "Vision — Evidence & Investigation Agents", "Agentic AI · investigation",
    "Investigation is manual: an adjuster hand-collects evidence and can't see cross-claim relationships, so complex fraud slips through.",
    "Autonomous agents that gather evidence, assemble a provenance-checked case file, and run network analysis — always handing a human the final call.",
    ["Retrieval + tool-use agents", "Graph analytics — relationship mining",
     "Provenance & chain-of-custody logic", "LLM — case narrative synthesis"],
    "Trigger → collect evidence → verify provenance → analyse network → case file → human decision.",
    "Investigator productivity multiplied; consistent, auditable case files.",
    "Seat-based investigator SaaS + per-case pricing.")

# 37 — Decision Intelligence + governance
feature_slide(37, "Vision — Decision Intelligence, Learning & Governance", "Decision & MLOps",
    "Detection without a governed decision layer means inconsistent outcomes and no way to learn safely or prove compliance.",
    "A decision engine that fuses all signals into an explained, policy-aware verdict, with continuous & federated learning, model monitoring and AI governance around it.",
    ["Ensemble decision fusion", "Continuous learning pipeline",
     "Federated learning — privacy-preserving", "Model monitoring & drift detection"],
    "Signals → policy-aware decision → explanation + audit → human-in-the-loop → feedback → retrain.",
    "Consistent, compliant, improving decisions across every insurer on the platform.",
    "Core platform tier — the layer everything else feeds.")

# 38 — Foundation model
feature_slide(38, "Vision — Insurance Foundation Model & Knowledge Graph", "Frontier · moat",
    "General LLMs don't understand insurance deeply; knowledge is fragmented across policies, claims and regulation.",
    "A private, insurance-specific foundation model plus an Insurance Knowledge Graph powering an ‘Insurance Copilot' for staff and customers.",
    ["Domain-adapted foundation model", "Insurance Knowledge Graph",
     "Retrieval-augmented Copilot", "Digital Twin & predictive fraud simulation"],
    "Curated insurance corpus → domain model + graph → Copilot & predictive services across the platform.",
    "A compounding data moat and the industry's default insurance intelligence.",
    "Platform-wide differentiator; licensing & premium intelligence tiers.",
    "This is the long-horizon vision — it depends on data accumulated by the earlier modules.")

# 39 — ML PORTFOLIO
s = new_slide()
header(s, "The Machine-Learning Portfolio", "Right model for each job — sequenced by roadmap.", section="Vision", kicker="ML / DL / CV")
rows = [
    ("Model / algorithm", "Primary use", "Phase"),
    ("Isolation Forest · Autoencoder", "Claim anomaly detection", "Fraud"),
    ("XGBoost · LightGBM · Random Forest", "Supervised fraud/risk scoring", "Fraud"),
    ("Graph Neural Networks", "Fraud-ring / network detection", "Fraud"),
    ("YOLO · SAM · CNN", "Damage detect / segment / grade", "Damage"),
    ("Vision Transformers", "OCR, forgery, severity", "OCR / Damage"),
    ("LSTM · Transformers", "Sequence & behaviour modelling", "Behaviour"),
    ("LLMs (provider-agnostic)", "Advisory, reasoning, Copilot", "LIVE + Vision"),
]
from_x, from_y = Inches(0.95), Inches(2.0); cw = [Inches(5.0), Inches(4.3), Inches(2.2)]; rh = Inches(0.56)
for r, row in enumerate(rows):
    x = from_x
    for c, cell in enumerate(row):
        head = (r == 0)
        fill = NAVY3 if head else (NAVY2 if r % 2 else NAVY)
        cellsp = rect(s, x, from_y + r * rh, cw[c], rh, fill, line=LINE)
        tf = cellsp.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.15); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = para(tf, first=True, space_after=0)
        col = CYAN if head else WHITE
        if c == 2 and not head:
            col = GREEN if "LIVE" in cell else AMBER
        run(p, cell, 12 if not head else 12.5, col, bold=head or c == 2)
        x += cw[c]
notes(s, N("Demonstrate ML breadth and that model choices are deliberate, not buzzwords.",
    "We're not sprinkling buzzwords. Each model earns its place: tree ensembles for tabular "
    "fraud risk, graph nets for rings, YOLO and SAM for damage, transformers for documents, "
    "and provider-agnostic LLMs — which are live — for reasoning and the future Copilot.",
    [("Why so many models?",
      "Because insurance spans tabular, image, text and graph data. One model can't cover it; "
      "the platform's job is to orchestrate the right one per task.")]))

# =============================================================================
#  SECTION G — ML / API / BUSINESS
# =============================================================================

# 40 — API PLATFORM
s = new_slide()
header(s, "The API Platform — How Insurers Integrate", "One integration, many intelligences.", section="Business", kicker="API-as-a-Service")
two_col(s,
    [('H', 'Integration surface'),
     ('B', 'REST APIs + SDKs for each intelligence module'),
     ('B', 'API keys, scoped auth, rate limits & quotas'),
     ('B', 'Developer portal: docs, sandbox, usage & billing'),
     ('B', 'Webhooks & async jobs for heavy CV workloads')],
    [('H', 'Enterprise integration'),
     ('B', 'Sits beside core systems — no rip-and-replace'),
     ('B', 'Multi-tenant isolation & per-tenant governance'),
     ('B', 'Explainability & audit exposed through the API'),
     ('B', 'Status: service API live; full platform is roadmap')])
caption(s, "API architecture: insurer app → API gateway (authz/rate-limit) → intelligence services → audit.")
notes(s, N("Explain the business-critical integration model and its current status.",
    "The go-to-market is API-as-a-Service. An insurer integrates once and calls whichever "
    "intelligences they need, priced by usage. A service API exists today; the full "
    "multi-tenant developer platform with keys, quotas and billing is the roadmap we're "
    "raising to build."))

# 41 — BUSINESS MODEL
s = new_slide()
header(s, "Business Model", "Multiple, compounding revenue lines off one platform.", section="Business", kicker="How we make money")
model = [("Enterprise SaaS", "Per-seat consoles for claims & investigation teams"),
         ("API usage pricing", "Per-call / per-inspection / per-verification"),
         ("Enterprise licence", "Annual platform contracts with SLAs"),
         ("White-label", "Insurer-branded advisory & portals"),
         ("Developer APIs", "Self-serve tiers for insurtech builders"),
         ("Public sector & health", "Govt schemes, hospitals, TPAs")]
gx, gy, cw, ch, mx, my = Inches(0.95), Inches(2.15), Inches(3.75), Inches(1.5), Inches(0.18), Inches(0.25)
for i, (t, d) in enumerate(model):
    col, row = i % 3, i // 3
    x = gx + col * (cw + mx); y = gy + row * (ch + my)
    pnl = panel(s, x, y, cw, ch, rgb=NAVY3)
    rect(s, x, y, cw, Inches(0.07), ROYAL)
    tf = pnl.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.16); tf.margin_right = Inches(0.15)
    p = para(tf, first=True, space_after=4); run(p, t, 14.5, WHITE, bold=True, font=FONT_SB)
    p2 = para(tf, space_after=0, line=1.05); run(p2, d, 11.5, MUTE)
notes(s, N("Show diversified, platform-driven revenue.",
    "Because it's a platform, revenue compounds: SaaS seats, usage-based API pricing, "
    "enterprise licences, white-label, self-serve developer tiers, and public-sector and "
    "health channels. Land with advisory, expand into the high-value fraud and claims APIs."))

# 42 — MARKET
s = new_slide()
header(s, "Market Opportunity", "A large, digitising, trust-starved industry.", section="Business", kicker="Why now")
two_col(s,
    [('H', 'The opening'),
     ('B', 'Global insurance is a multi-trillion-dollar premium market'),
     ('B', 'Fraud leakage is widely estimated in the hundreds of billions/yr'),
     ('B', 'Insurers are actively buying AI but distrust black boxes'),
     ('B2', 'Figures are cited industry estimates, not our measurements')],
    [('H', 'Why now'),
     ('B', 'LLMs made explainable, conversational AI viable'),
     ('B', 'CV & graph ML matured for fraud/damage'),
     ('B', 'Regulation is demanding transparency — our default'),
     ('B', 'API-first integration lowers the adoption barrier')])
caption(s, "TAM/SAM/SOM funnel (labelled ‘illustrative, from cited industry estimates').")
notes(s, N("Give a credible ‘why now' without inventing precise numbers.",
    "The market is enormous and digitising, and — critically — insurers are buying AI but "
    "don't trust it. The timing is right because LLMs, computer vision and graph ML have all "
    "matured, and regulation now rewards exactly the transparency we build in. I cite market "
    "figures as industry estimates, clearly labelled.",
    [("What's your TAM?",
      "Global insurance premiums are multi-trillion; even a small AI-services slice is a "
      "large market. I present these as cited third-party estimates, and our beachhead is "
      "advisory + fraud APIs, not the whole TAM.")]))

# 43 — WHY NOT NORMAL PROJECT
s = new_slide()
header(s, "Why This Is Not a Normal College Project", section="Positioning", kicker="Altitude")
two_col(s,
    [('H', 'What it is NOT'),
     ('B', 'Not a single chatbot or a toy demo'),
     ('B', 'Not one model doing one narrow task'),
     ('B', 'Not throwaway coursework code')],
    [('H', 'What it IS'),
     ('B', 'An enterprise AI platform with real architecture'),
     ('B', 'AI infrastructure delivered as APIs'),
     ('B', 'A startup product with a business model'),
     ('B', 'A research project mapping gaps → build'),
     ('B', 'A path to a private insurance foundation model')])
caption(s, "Altitude ladder: script → app → product → platform → infrastructure → foundation model.")
notes(s, N("Directly claim the altitude — respectfully but firmly.",
    "I want to be direct about ambition. This is engineered as a platform and a company, not "
    "a course submission. It has real architecture, tests, security, a business model and a "
    "research foundation. That's the bar we're holding ourselves to."))

# 44 — ROADMAP (drawn timeline)
s = new_slide()
header(s, "Product Roadmap", "Five phases from coursework to category-defining infrastructure.", section="Roadmap", kicker="The 5-phase plan")
phases = [("Phase 1", "College Project", "Multi-agent advisory MVP — DONE", GREEN),
          ("Phase 2", "Startup MVP", "Fraud · OCR · Damage v1 + API", AMBER),
          ("Phase 3", "Enterprise SaaS", "Multi-tenant, governance, console", AMBER),
          ("Phase 4", "Global API Platform", "Developer ecosystem & marketplace", ROYAL),
          ("Phase 5", "Foundation Model", "Insurance model + knowledge graph", CYAN)]
n = len(phases); tw = Inches(2.3); gap = Inches(0.28); startx = Inches(0.95); y = Inches(3.0)
# baseline
rect(s, startx, y + Inches(1.05), Inches(11.4), Inches(0.05), LINE)
for i, (ph, title, d, c) in enumerate(phases):
    x = startx + i * (tw + gap)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x + tw / 2 - Inches(0.12), y + Inches(0.95), Inches(0.24), Inches(0.24))
    solid(dot, c)
    pnl = panel(s, x, y, tw, Inches(0.85), rgb=NAVY3)
    rect(s, x, y, tw, Inches(0.07), c)
    tf = pnl.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.14); tf.margin_top = Inches(0.1)
    p = para(tf, first=True, space_after=1); run(p, ph, 11, c, bold=True)
    p2 = para(tf, space_after=0); run(p2, title, 13, WHITE, bold=True, font=FONT_SB)
    _, tf2 = textbox(s, x, y + Inches(1.45), tw, Inches(1.3))
    pp = para(tf2, first=True, space_after=0, line=1.05); run(pp, d, 11, MUTE)
notes(s, N("Anchor everything on a concrete, staged plan with Phase 1 already done.",
    "Phase 1 — the advisory platform — is done. Phase 2 adds the first CV and fraud modules "
    "and the API. Phases 3 to 5 scale it to multi-tenant SaaS, a global API platform, and "
    "finally an insurance foundation model. Funding accelerates Phase 2 to 3.",
    [("What would investment buy?",
      "Phase 2 execution: the fraud/OCR/damage APIs, a design-partner insurer, and the "
      "multi-tenant foundations for Phase 3.")]))

# 45 — DIFFERENTIATION
s = new_slide()
header(s, "Competitive Differentiation", "Everyone sells a feature; we sell the decision layer.", section="Positioning", kicker="Moat")
rows = [
    ("Dimension", "Point solutions / insurtech", "Aegis AI"),
    ("Scope", "One feature (OCR, or a bot)", "Composable intelligence platform"),
    ("Decisions", "Scores only", "Explained, evidence-based decisions"),
    ("Trust", "Black box", "Explainable + human-in-the-loop + audit"),
    ("Integration", "Bespoke", "API-first, augments core systems"),
    ("Moat", "Model of the month", "Data → knowledge graph → foundation model"),
]
from_x, from_y = Inches(0.95), Inches(2.05); cw = [Inches(2.8), Inches(4.3), Inches(4.3)]; rh = Inches(0.66)
for r, row in enumerate(rows):
    x = from_x
    for c, cell in enumerate(row):
        head = (r == 0)
        fill = NAVY3 if head else (NAVY2 if r % 2 else NAVY)
        cellsp = rect(s, x, from_y + r * rh, cw[c], rh, fill, line=LINE)
        tf = cellsp.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.15); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = para(tf, first=True, space_after=0)
        col = CYAN if head else (GREEN if c == 2 else MUTE)
        run(p, cell, 12.5 if head else 12, col, bold=head or c == 2)
        x += cw[c]
notes(s, N("Frame the moat as the decision layer + data compounding.",
    "Competitors ship features and get commoditised. Our defensibility is the decision and "
    "trust layer on top, and the data that compounds through it into a knowledge graph and "
    "eventually a foundation model. That's a widening moat, not a model of the month."))

# 46 — IMPACT
s = new_slide()
header(s, "Expected Impact", section="Positioning", kicker="Outcomes")
two_col(s,
    [('H', 'For insurers'),
     ('B', 'Lower fraud leakage & loss-adjustment cost'),
     ('B', 'Faster, consistent, defensible claims decisions'),
     ('B', 'One governed AI layer instead of many point tools')],
    [('H', 'For customers & society'),
     ('B', 'Claims in minutes, not weeks'),
     ('B', 'Fairer pricing as leakage falls'),
     ('B', 'Accessible, multilingual, trustworthy guidance'),
     ('B', 'AI decisions people can actually understand')])
caption(s, "Impact dashboard mock: leakage↓, cycle-time↓, CSAT↑, explainability-coverage↑.")
notes(s, N("Close the problem loop: impact for insurer, customer and society.",
    "The impact is symmetric: insurers save money and defend decisions; customers get speed, "
    "fairness and understanding. And because we reach underserved users, there's a genuine "
    "social dimension — insurance that finally explains itself."))

# 47 — TRACTION / STATUS
s = new_slide()
header(s, "Where We Are Today", "Honest status — a working platform, ready to extend.", section="Traction", kicker="Current status")
stats = [("5", "Live specialist AI agents"), ("36", "Curated plans, 4 categories"),
         ("200+", "Automated tests + CI"), ("3", "Provider-agnostic LLM backends"),
         ("2", "Live portals (customer/admin)"), ("Phase 1", "Roadmap complete")]
gx, gy, cw, ch, mx, my = Inches(0.95), Inches(2.2), Inches(3.75), Inches(1.5), Inches(0.18), Inches(0.25)
for i, (v, d) in enumerate(stats):
    col, row = i % 3, i // 3
    x = gx + col * (cw + mx); y = gy + row * (ch + my)
    pnl = panel(s, x, y, cw, ch, rgb=NAVY3)
    tf = pnl.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = para(tf, first=True, space_after=2, align=PP_ALIGN.CENTER)
    run(p, v, 30, CYAN, bold=True, font=FONT_SB)
    p2 = para(tf, space_after=0, align=PP_ALIGN.CENTER); run(p2, d, 12, MUTE)
notes(s, N("Ground the pitch in verifiable, current metrics — all from the repo.",
    "These numbers are all verifiable in the codebase, not marketing. Five live agents, a "
    "36-plan catalogue, 200-plus tests, provider-agnostic LLMs, two live portals, Phase 1 "
    "complete. We're past the risky ‘does anything work' stage.",
    [("Any customers or revenue yet?",
      "Not yet — this is a final-year build seeking its first design-partner insurer. The ask "
      "is a pilot, not a claim of traction we don't have.")]))

# 48 — JURY Q&A
s = new_slide()
header(s, "Anticipated Questions", "Straight answers, prepared.", section="Q&A", kicker="For the panel")
qa = [
    ("‘Isn't this just a chatbot?'", "No — it's a multi-agent platform with orchestration, memory, recommendation and an API layer; the bot is one surface."),
    ("‘Do the CV / fraud models exist?'", "Not yet — clearly labelled roadmap. The platform that hosts them is built and tested."),
    ("‘How is it explainable?'", "Live reasoning trace + reason-for-fit today; evidence-linked decision receipts on the roadmap."),
    ("‘Why will insurers adopt?'", "API-first: augment, don't replace. Weeks to integrate, governed and auditable by design."),
    ("‘What's the moat?'", "The decision/trust layer plus compounding data → knowledge graph → foundation model."),
    ("‘Where are the market numbers from?'", "Cited third-party industry estimates, labelled as estimates — never our own measurements."),
]
for i, (q, a) in enumerate(qa):
    y = Inches(2.1) + i * Inches(0.78)
    pnl = panel(s, Inches(0.95), y, Inches(11.4), Inches(0.68))
    tf = pnl.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    p = para(tf, first=True, space_after=0, line=1.0)
    run(p, q + "   ", 13, CYAN, bold=True)
    run(p, a, 12, WHITE)
notes(s, N("Pre-empt the hard questions; answer with the honesty theme.",
    "I'll leave these here for the panel. The through-line in every answer is the same: be "
    "precise about built vs. roadmap, lead with trust, and let the working system speak. "
    "Happy to go deep on any of them.",
    [("What's the single biggest risk?",
      "Execution risk on the CV/fraud models and landing a first insurer. We de-risk by "
      "building on a proven platform and starting with a narrow, high-value fraud/OCR wedge.")]))

# =============================================================================
#  SECTION H — CLOSE
# =============================================================================

# 49 — FUTURE VISION
s = new_slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, Inches(0.9), Inches(2.0), Inches(1.4), Inches(0.11), CYAN)
_, tf = textbox(s, Inches(0.9), Inches(2.25), Inches(11.5), Inches(3))
p = para(tf, first=True, space_after=0)
run(p, "The Future We're Building", 40, WHITE, bold=True, font=FONT_SB)
p2 = para(tf, space_before=14, space_after=0, line=1.2)
run(p2, "Every insurer on Earth running trustworthy AI — fraud stopped with evidence, "
        "claims settled in minutes, and every decision one a human can understand.", 21, MUTE)
p3 = para(tf, space_before=18, space_after=0)
run(p3, "Aegis AI — the AI operating system for the global insurance industry.", 18, CYAN, bold=True)
footer(s, "Vision")
notes(s, N("Lift back up to the vision before the close.",
    "Zoom back out. The world we're building is one where insurance AI is trusted by "
    "default — fraud stopped with evidence, claims in minutes, every decision explainable. "
    "Aegis is the operating system that gets the industry there."))

# 50 — REFERENCES
s = new_slide()
header(s, "References", "IEEE style — selected.", section="References", kicker="Selected references")
refs = [
    "[1] “AI Revolution in Insurance: Bridging Research and Reality,” (anchor paper for this work).",
    "[2] F. T. Liu et al., “Isolation Forest,” IEEE ICDM, 2008.",
    "[3] T. Chen and C. Guestrin, “XGBoost: A Scalable Tree Boosting System,” ACM KDD, 2016.",
    "[4] T. N. Kipf and M. Welling, “Semi-Supervised Classification with Graph Convolutional Networks,” ICLR, 2017.",
    "[5] J. Redmon et al., “You Only Look Once (YOLO): Unified, Real-Time Object Detection,” IEEE CVPR, 2016.",
    "[6] A. Kirillov et al., “Segment Anything (SAM),” IEEE/CVF ICCV, 2023.",
    "[7] A. Dosovitskiy et al., “An Image is Worth 16×16 Words: Transformers for Image Recognition,” ICLR, 2021.",
    "[8] A. Vaswani et al., “Attention Is All You Need,” NeurIPS, 2017.",
    "[9] S. Lundberg and S.-I. Lee, “A Unified Approach to Interpreting Model Predictions (SHAP),” NeurIPS, 2017.",
    "[10] B. McMahan et al., “Communication-Efficient Learning of Deep Networks from Decentralized Data (Federated Learning),” AISTATS, 2017.",
    "[11] Industry fraud-cost estimates: Coalition Against Insurance Fraud; Swiss Re / Deloitte insurance-AI reports.",
    "[12] Technical reports & model documentation: OpenAI, Google, Anthropic, Microsoft Research, NVIDIA.",
]
_, tf = textbox(s, Inches(0.95), Inches(1.95), Inches(11.5), Inches(4.8))
for i, r in enumerate(refs):
    p = para(tf, first=(i == 0), space_after=6, line=1.02)
    run(p, r, 12, MUTE)
notes(s, N("Signal scholarship and honest sourcing.",
    "References are IEEE-style and real: the anchor paper, the core ML/CV/graph and "
    "explainability works, federated learning, and vendor technical reports. Market figures "
    "are attributed to cited industry sources, consistent with how I presented them."))

# 51 — THANK YOU / CONTACT / CLOSE
s = new_slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, Inches(0.9), Inches(1.75), Inches(1.4), Inches(0.11), CYAN)
_, tf = textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2))
p = para(tf, first=True, space_after=0)
run(p, "Thank You", 54, WHITE, bold=True, font=FONT_SB)
p2 = para(tf, space_before=10, space_after=0)
run(p2, "Let's build the AI operating system for insurance — together.", 20, CYAN, bold=True)
# investor closing
panel(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.25))
_, tfc = textbox(s, Inches(1.2), Inches(4.2), Inches(11.0), Inches(0.9), anchor=MSO_ANCHOR.MIDDLE)
pc = para(tfc, first=True, space_after=0, line=1.15)
run(pc, "We've built the platform and proven the architecture. The ask: a design-partner "
        "insurer and the backing to ship Phase 2 — fraud, OCR and damage intelligence — as "
        "the first APIs of an industry standard.", 15, WHITE)
_, tf3 = textbox(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1.0))
p3 = para(tf3, first=True, space_after=2)
run(p3, "Sri Nivaspanneer", 16, WHITE, bold=True)
p4 = para(tf3, space_after=0)
run(p4, "srinivaspanneerlm@gmail.com     ·     Aegis AI — The AI Operating System for Insurance", 13, MUTE)
footer(s, "Close")
notes(s, N("End on the ask, not just thanks.",
    "So to close: we've built the platform and proven the architecture. What I'm looking for "
    "is a design-partner insurer and the backing to ship Phase 2 — the fraud, OCR and damage "
    "APIs. That's how Aegis becomes the industry's default intelligence layer. Thank you — "
    "I'd love your questions."))

# ── Save ─────────────────────────────────────────────────────────────────────
out = "AegisAI_Pitch.pptx"
prs.save(out)
print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides")
